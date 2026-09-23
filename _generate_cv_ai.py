from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)


def set_run_font(run, name="Calibri", size=10, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_horizontal_line(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    pBdr.append(bottom)
    pPr.append(pBdr)


def section_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text.upper())
    set_run_font(run, size=11, bold=True, color=(31, 78, 121))
    add_horizontal_line(p)
    return p


def body_para(text, space_before=2, space_after=2, size=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_run_font(run, size=size)
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.left_indent = Inches(0.2)
    run = p.add_run(text)
    set_run_font(run, size=10)
    return p


def role_header(title, dates):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run1 = p.add_run(title)
    set_run_font(run1, size=10, bold=True)
    run2 = p.add_run("\t" + dates)
    set_run_font(run2, size=10, bold=False, color=(80, 80, 80))
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(7.0), alignment=2)
    return p


def company_line(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_run_font(run, size=10, bold=False, color=(70, 70, 70))
    run.italic = True
    return p


# ========== HEADER ==========
name = doc.add_paragraph()
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
name.paragraph_format.space_before = Pt(0)
name.paragraph_format.space_after = Pt(2)
r = name.add_run("SARTHAK KULKARNI")
set_run_font(r, size=18, bold=True, color=(31, 78, 121))

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(2)
r = title.add_run("GenAI & DevOps Engineer  |  Agentic AI Systems")
set_run_font(r, size=11, bold=True, color=(60, 60, 60))

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.paragraph_format.space_before = Pt(0)
contact.paragraph_format.space_after = Pt(2)
r = contact.add_run(
    "Pune, Maharashtra, India  |  (+91) 93256-02791  |  sarthakkul2311@gmail.com\n"
    "linkedin.com/in/sarthak-kulkarni-/"
)
set_run_font(r, size=9.5, color=(70, 70, 70))

# ========== PROFILE ==========
section_heading("Professional Summary")
body_para(
    "GenAI and DevOps Engineer with 2+ years of experience designing and deploying production-grade "
    "agentic AI systems on Microsoft Azure. Specialized in multi-agent architectures, RAG pipelines, "
    "and Azure AI Foundry-based solutions that automate complex enterprise workflows. Proven track record "
    "of reducing manual effort through intelligent agents, building end-to-end CI/CD pipelines, and "
    "delivering scalable applications across Azure Web Apps, Function Apps, and Azure DevOps. Strong "
    "foundation in Python, cloud engineering, and full-stack development, with demonstrated leadership "
    "in mentoring teams and driving innovation."
)

# ========== SKILLS ==========
section_heading("Technical Skills")

skills = [
    (
        "Agentic AI & GenAI: ",
        "Multi-Agent Systems, Azure AI Foundry, RAG (Azure AI Search), Prompt Engineering, "
        "LLM Orchestration, Agentic Workflows",
    ),
    (
        "Cloud & DevOps: ",
        "Microsoft Azure, Azure DevOps, CI/CD Pipelines, Azure Web Apps, Azure Function Apps, "
        "Infrastructure Deployment",
    ),
    (
        "Data & Analytics: ",
        "Python, SQL, Databricks, PySpark, Azure Data Factory, Machine Learning, Time Series Analysis",
    ),
    (
        "Development: ",
        "Python, C++, Java, JavaScript, ReactJS, HTML/CSS, REST APIs, Git",
    ),
    (
        "Certifications Stack: ",
        "AZ-204 (Azure Developer Associate), AI-900, AZ-900, PCEP (Python)",
    ),
]

for label, content in skills:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    rb = p.add_run(label)
    set_run_font(rb, size=10, bold=True)
    rc = p.add_run(content)
    set_run_font(rc, size=10)

# ========== EXPERIENCE ==========
section_heading("Professional Experience")

role_header("GenAI Developer / DevOps Engineer", "August 2024 – Present")
company_line("Hexaware Technologies, Pune, India")

bullet(
    "Designed and deployed a production triaging agent on Azure that autonomously identifies issues, "
    "logs tickets, and resolves in-scope incidents—reducing manual operational effort by approximately 70%."
)
bullet(
    "Architected a multi-agent application to generate Informed Consent Forms (ICF) across N countries × N studies, "
    "orchestrating country-specific regulatory documents, customization plans, and review checklists via Azure AI Foundry."
)
bullet(
    "Built end-to-end agentic solutions using RAG (Azure AI Search), Azure Web Apps, Azure Function Apps, "
    "and Azure DevOps CI/CD pipelines for reliable cloud deployment and release automation."
)
bullet(
    "Mentored and led a cohort of 5 interns through the design and delivery of a competitive, "
    "production-oriented GenAI project."
)
bullet(
    "Won the Green Eco-fuelers Hackathon by delivering a sustainable solution powered by agentic AI systems."
)
bullet(
    "Won the Innovation Pitch competition by presenting a business concept and delivering a working proof-of-concept."
)

role_header("Front-End Web Developer", "July 2023 – September 2023")
company_line("Growdigis IT Solution Pvt. Ltd., Pune, India")
bullet(
    "Developed responsive web interfaces with backend integration for client-facing digital service platforms."
)
bullet(
    "Deployed a production website enabling a service provider to digitize and scale their business operations."
)

role_header("Data Science Intern", "April 2023 – May 2023")
company_line("CodeClause, Pune, India")
bullet(
    "Built a customer segmentation model using K-means clustering to support trend analysis and targeted insights."
)
bullet(
    "Developed an AI-based age and gender detection model, covering data preprocessing, training, and evaluation."
)

role_header("Trainee Software Developer", "September 2021 – December 2021")
company_line("White Code Technology Solutions Pvt. Ltd., Pune, India")
bullet(
    "Developed responsive websites for educational institutions with clean, accessible UI patterns."
)
bullet(
    "Researched and prototyped smart contracts and NFT creation workflows on blockchain platforms."
)

role_header("Growth Ambassador", "June 2021 – July 2021")
company_line("Younity.in, Delhi, India")
bullet(
    "Led a team of 15 interns, coordinating task allocation, development strategy, and delivery timelines."
)
bullet(
    "Supported business development and sales coordination initiatives across growth programs."
)

# ========== PROJECTS ==========
section_heading("Key Projects")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(3)
p.paragraph_format.space_after = Pt(1)
r = p.add_run("Sitepay Helpdesk Agent (Agentic Triaging System)")
set_run_font(r, size=10, bold=True)

bullet(
    "Engineered an autonomous helpdesk agent that ingests emails, classifies intent, logs tickets in "
    "Assyst/ServiceNow, resolves in-scope issues, closes tickets, and replies to clients—fully automating "
    "the QA and support workflow."
)
bullet(
    "Stack: Azure AI Foundry (Agents), Azure AI Search (RAG), Azure Web Apps, Azure Function Apps, Azure DevOps CI/CD."
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(1)
r = p.add_run("ICF Generator (Multi-Agent Document Automation)")
set_run_font(r, size=10, bold=True)

bullet(
    "Built a multi-agent system that generates country- and study-specific Informed Consent Forms by incorporating "
    "regulatory documents, customization plans, and review checklists—compressing a ~30-day manual process to "
    "approximately 10 minutes."
)
bullet(
    "Stack: Azure AI Foundry (Agents), Azure AI Search (RAG), Azure Web Apps, Azure Function Apps, Azure DevOps CI/CD."
)

# ========== EDUCATION ==========
section_heading("Education")

role_header(
    "Bachelor of Computer Science Engineering (Honors in Data Science)",
    "May 2024",
)
company_line(
    "Zeal College of Engineering & Research — Savitribai Phule Pune University, Pune"
)

role_header("Higher Secondary (Bifocal – Computer Science)", "May 2020")
company_line("MTES Jr. College — Maharashtra State Board, Pune")

role_header("Secondary Education", "April 2018")
company_line(
    "Dnyanganga English Medium School — Cyber & Science Olympiad; Hindi Rashtrabhasha"
)

# ========== PUBLICATIONS ==========
section_heading("Publications & Research")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(0)
r = p.add_run("DizBoard: A Monitoring Dashboard Using Streamlit")
set_run_font(r, size=10, bold=True)
r2 = p.add_run("  |  IJSREM")
set_run_font(r2, size=10, color=(70, 70, 70))
bullet(
    "Developed an interactive Streamlit dashboard for real-time data visualization and operational analytics."
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(3)
p.paragraph_format.space_after = Pt(0)
r = p.add_run(
    "A Survey on Machine Learning Algorithms for Risk-Controlled Algorithmic Trading"
)
set_run_font(r, size=10, bold=True)
r2 = p.add_run("  |  IJSRST")
set_run_font(r2, size=10, color=(70, 70, 70))
bullet(
    "Surveyed machine learning techniques applicable to risk-controlled algorithmic trading strategies."
)

# ========== CERTIFICATIONS ==========
section_heading("Certifications")
for c in [
    "AZ-204: Microsoft Azure Developer Associate",
    "AI-900: Microsoft Azure AI Fundamentals",
    "AZ-900: Microsoft Azure Fundamentals",
    "PCEP: Python Certified Entry-Level Programmer",
]:
    bullet(c)

# ========== LEADERSHIP ==========
section_heading("Leadership & Extracurricular")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(1)
p.paragraph_format.space_after = Pt(0)
rb = p.add_run("President, Association of Computer Engineering Students — ")
set_run_font(rb, size=10, bold=True)
rc = p.add_run("Organized university-level hackathons and technical/cultural events.")
set_run_font(rc, size=10)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(1)
p.paragraph_format.space_after = Pt(0)
rb = p.add_run("Student Sports Coordinator, Zeal Cricket Club — ")
set_run_font(rb, size=10, bold=True)
rc = p.add_run(
    "Led the college cricket team; managed tournament logistics and coordination."
)
set_run_font(rc, size=10)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(1)
p.paragraph_format.space_after = Pt(0)
rb = p.add_run("Acting Club Coordinator, Zeal Cultural Center — ")
set_run_font(rb, size=10, bold=True)
rc = p.add_run("Directed and performed in state-level drama competitions.")
set_run_font(rc, size=10)

# ========== LANGUAGES ==========
section_heading("Languages")
body_para(
    "English (Fluent)  |  Hindi (Fluent)  |  Marathi (Native)  |  German (Basic)"
)

out_path = r"c:\Users\KulkarS4.EU\Documents\portfolio-web\portfolio\Sarthak-CV-AI.docx"
doc.save(out_path)
print("Saved:", out_path)
